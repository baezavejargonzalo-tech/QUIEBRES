"""Forecast propuesto de una semana por SKU y cadena, para comparar con tu FCST.

Uso:
    python3 forecast_mes1/forecast_semana.py FCST.xlsx --semana 44 --anio 2026 \
        --venta-real tp.xlsx --out forecast_mes1/output/forecast_S44.xlsx

Por SKU×cadena con historia en la hoja "Precio vs SellIn 25-26":
  A = Sell In del año anterior en la semana (promedio centrado de 3 semanas) × ritmo año actual / anterior
  B = ritmo reciente del año actual × estacionalidad (semana objetivo / semanas recientes, año anterior);
      el ritmo reciente pesa más las últimas semanas (el peso se reduce a la mitad cada 3 semanas)
  base = promedio de A y B. Ritmo y estacionalidad se suavizan hacia los de la cadena.
  Tottus con promo: × curva de la hoja Leyenda según la semana relativa al inicio.
Canal Tradicional: venta normal real de las semanas del archivo TP (promedio semanal).
Gerencias o cadenas sin historia (CCU, JUNAEB, INDUSTRIAL, SKU nuevos): se mantiene tu FCST.
El precio ya está en la base (usa venta del año actual a precios actuales).
"""

import argparse
import datetime as dt
import json
import re
import statistics
from collections import defaultdict
from pathlib import Path

import openpyxl

from build_forecast_mes1 import leer_fcst, leer_precios, leer_venta_real, num

K_YOY, K_ESTAC = 30.0, 2.0
YOY_MIN, YOY_MAX = 0.3, 3.0
ESTAC_MIN, ESTAC_MAX = 0.6, 1.6
SEM_RR = 8
VIDA_MEDIA = 3        # semanas: el peso de cada semana del ritmo reciente se reduce a la mitad cada 3 semanas hacia atrás
CURVA_TOTTUS = {-2: 0.64, -1: 1.69, 0: 4.32, 1: 2.48, 2: 2.73, 3: 1.25, 4: 0.46}


# Datos que más mejorarían la semana (se muestran en la página)
FALTAN = [
    ("1 · Reporte CPFR al cierre de la S42 y la S43", "Cuando cierren, el mismo reporte de semana 40–44 deja el ritmo reciente y el sesgo de tu forecast al día."),
    ("2 · Órdenes de compra ya recibidas para la S44", "Por cadena y SKU (Walmart, Cencosud, Tottus, Unimarc). Es lo que más acota la semana: lo ya pedido es piso del forecast."),
    ("3 · Stock y Sell Out en tienda", "Días de inventario por SKU en Walmart (Retail Link) y Cencosud (B2B). Si la cadena está sobrestockeada, compra menos aunque venda igual."),
    ("4 · Fecha de inicio de la promo Cencosud de noviembre", "Si parte el lunes 26 de octubre, parte del volumen pactado cae en la S44. Hoy asumo que parte en noviembre y no la sumo a la S44."),
    ("5 · Pedidos de Operadores Comerciales", "Compran néctares sixpack semana por medio en pedidos grandes (S36, S38 y S41 altos; S37, S39 y S40 bajos). Si confirmas si la S44 es semana de pedido, Canal Tradicional deja de ser un promedio."),
    ("6 · Catálogo Tottus de noviembre", "La promo Tottus de la S40 cae a 0,46× en la S44 según tu curva. Si hay nueva actividad desde la S44, cambia el número."),
]


# cadena del reporte CPFR para cada subcadena del forecast
CPFR_CADENA = {"WALMART": "Walmart", "JUMBO": "Cencosud", "SANTA ISABEL": "Cencosud", "UNIMARC": "Unimarc",
               "TOTTUS": "Tottus", "ALVI": "Alvi", "SUPERMERCADOS REGION": "Supermercados Region",
               "CANAL TRADICIONAL": "Canal Tradicional", "CCU · ECUSA": "Ccu",
               "INDUSTRIAL Y FOOD SERVICE": "Industrial Y Food Service",
               "INDUSTRIAL Y FOOD SERVICE · JUNAEB": "Industrial Y Food Service",
               "INDUSTRIAL Y FOOD SERVICE · INDUSTRIAL": "Industrial Y Food Service"}


def leer_cpfr(path, skus):
    """Reporte CPFR semanal (FCST, Solicitado, Venta Sell IN, Quebrados por SKU y cadena).
    Semanas cerradas = las que ya tienen Sell In; las abiertas solo informan pedidos ya recibidos."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    filas = wb.worksheets[0].iter_rows(values_only=True)
    h = {str(c).strip(): i for i, c in enumerate(next(filas)) if c is not None}
    acc = defaultdict(lambda: defaultdict(float))
    for r in filas:
        if r[h["SKU"]] not in skus:
            continue
        w = int(str(r[h["Semana"]])[-2:])
        k = (w, r[h["CADENA"]], r[h["SKU"]])
        for c, n in (("FCST", "fcst"), ("Solicitado", "sol"), ("Venta Sell IN", "si"), ("Quebrados", "queb")):
            acc[k][n] += num(r[h[c]])
    por_sem = defaultdict(float)
    for (w, _, _), v in acc.items():
        por_sem[w] += v["si"]
    cerradas = sorted(w for w, v in por_sem.items() if v > 0)
    abiertas = sorted(w for w in por_sem if w not in cerradas)
    cad = defaultdict(lambda: defaultdict(float))
    sku_cad = defaultdict(lambda: defaultdict(float))
    ab = defaultdict(float)
    for (w, c, sku), v in acc.items():
        if w in cerradas:
            for n, x in v.items():
                cad[c][n] += x
                sku_cad[(sku, c)][n] += x
        else:
            ab[(w, c)] += v["sol"]
    return {"cerradas": cerradas, "abiertas": abiertas,
            "cadenas": {c: dict(v) for c, v in cad.items()},
            "sku_cad": {f"{sku}|{c}": dict(v) for (sku, c), v in sku_cad.items()},
            "pedidos_abiertas": [{"semana": w, "cadena": c, "sol": v} for (w, c), v in sorted(ab.items()) if v > 0.05]}


def leer_no_normal(path, semanas):
    """Venta intermedia y liquidación (TP) por cadena del CPFR y semana: el CPFR la cuenta como pedido,
    pero tu FCST no la incluye."""
    from build_forecast_mes1 import EXCLUIR_RUT_TRADICIONAL
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    filas = wb.worksheets[0].iter_rows(values_only=True)
    h = {str(c).strip(): i for i, c in enumerate(next(filas)) if c is not None}
    ger_cad = {"WALMART": "Walmart", "CENCOSUD": "Cencosud", "SMU": "Unimarc", "TOTTUS": "Tottus",
               "SUPERMERCADOS REGIONALES": "Supermercados Region", "INDUSTRIAL Y FOOD SERVICE": "Industrial Y Food Service"}
    out = defaultdict(float)
    for r in filas:
        tipo = r[h["Tipo de Venta"]]
        if tipo not in ("VENTA INTERMEDIA", "VENTA LIQUIDACION") or r[h["Semana"]] not in semanas:
            continue
        ger, cad, rut = r[h["Gerencia de Venta"]], r[h["Cadena Cliente"]], str(r[h.get("Rut Cliente", 0)]).strip()
        if rut in EXCLUIR_RUT_TRADICIONAL:
            c = "Ccu"
        elif ger == "TRADICIONAL Y REGIONES":
            c = "Alvi" if cad == "ALVI SUPERMERCADOS" else "Canal Tradicional"
        else:
            c = ger_cad.get(ger, ger)
        out[c] += num(r[h["Venta Fisica SelI In (TON)"]])
    return dict(out)


def inicio_promo(dur):
    if isinstance(dur, dt.datetime):
        return dur.date().isocalendar()[1]
    m = re.search(r"(\d{1,2})-(\d{1,2})-(\d{4})", str(dur or ""))
    return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))).isocalendar()[1] if m else None


def construir(xlsx, semana, anio, venta_real=None, cpfr=None, tp_path=None):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    fcst, semanas_fcst = leer_fcst(wb, 11)
    hist = leer_precios(wb, anio)
    if semana not in semanas_fcst:
        raise SystemExit(f"La semana S{semana} no está en la hoja FCST")
    reales = sorted({w for d in hist.values() for w, v in d.items() if v["si1"] is not None})
    ult = max(reales)
    sem_yoy = reales
    sem_rr = [w for w in reales if w > ult - SEM_RR]
    _p = {w: 0.5 ** ((ult - w) / VIDA_MEDIA) for w in sem_rr}
    pes = {w: v / sum(_p.values()) for w, v in _p.items()}
    vent = [semana - 1, semana, semana + 1]
    vr, vr_sem = venta_real if venta_real else ({}, [])
    vr_w = [w for _, w in vr_sem]

    g = lambda h, w, k: num(h.get(w, {}).get(k))

    def factores(h):
        b0 = sum(g(h, w, "si0") for w in sem_yoy)
        b1 = sum(g(h, w, "si1") for w in sem_yoy)
        rr0 = sum(pes[w] * g(h, w, "si0") for w in sem_rr)
        rr1 = sum(pes[w] * g(h, w, "si1") for w in sem_rr)
        s0 = sum(g(h, w, "si0") for w in vent) / len(vent)
        return b0, b1, rr0, rr1, s0

    # factores por cadena (prior)
    tot_cad = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
    for (sku, sub), h in hist.items():
        for w, v in h.items():
            for k in ("si0", "si1"):
                tot_cad[sub][w][k] += num(v[k])
    prior = {}
    for sub, h in tot_cad.items():
        b0, b1, rr0, rr1, s0 = factores(h)
        prior[sub] = {"yoy": b1 / b0 if b0 else 1.0, "estac": s0 / rr0 if rr0 else 1.0}

    filas = defaultdict(lambda: {"fcst": 0.0, "real_rec": [], "promos": []})
    for f in fcst:
        r = filas[(f["sku"], f["subcadena"])]
        r.update({"sku": f["sku"], "producto": f["producto"], "subcat": f["subcat"], "gerencia": f["gerencia"],
                  "cadena": f["subcadena"]})
        r["fcst"] += f["semanas"].get(semana, 0)
        r["fcst_prev"] = r.get("fcst_prev", 0) + sum(f["semanas"].get(w, 0) for w in (semana - 2, semana - 1)) / 2
        if f["promo"] and not re.match(r"^[A-Z]+\d{8}$", str(f["promo"])):
            r["promos"].append({"nombre": str(f["promo"]), "dur": f["promo_dur"], "vol": f["vol_pactado"]})

    out = []
    for (sku, sub), r in filas.items():
        h = hist.get((sku, sub))
        row = {k: r[k] for k in ("sku", "producto", "subcat", "gerencia", "cadena")}
        row.update({"fcst": r["fcst"], "fcst_prev": r["fcst_prev"], "promo": "; ".join(sorted({p["nombre"].split(" (")[0] for p in r["promos"]}))})
        if h:
            b0, b1, rr0, rr1, s0 = factores(h)
            pr = prior.get(sub, {"yoy": 1.0, "estac": 1.0})
            yoy = (b1 + K_YOY * pr["yoy"]) / (b0 + K_YOY) if b0 > 0 else None
            yoy = min(max(yoy, YOY_MIN), YOY_MAX) if yoy else None
            estac = (s0 + K_ESTAC * pr["estac"]) / (rr0 + K_ESTAC) if rr0 > 0 else pr["estac"]
            estac = min(max(estac, ESTAC_MIN), ESTAC_MAX)
            A = s0 * yoy if (yoy and s0 > 0) else None
            B = rr1 * estac if rr1 > 0 else None
            ms = [m for m in (A, B) if m is not None]
            base = sum(ms) / len(ms) if ms else 0.0
            sem_vals = [g(h, w, "si1") for w in sem_rr]
            media = statistics.mean(sem_vals)
            cv = statistics.pstdev(sem_vals) / media if media > 0 else None
            fuente = "historia"
            nuevo = base < 0.05 and r["fcst"] >= 0.5
            if nuevo:
                base, fuente = r["fcst"], "sin venta previa (tu FCST)"
            mult, nota = 1.0, ""
            for p in r["promos"]:
                if "tottus" in p["nombre"].lower():
                    w0 = inicio_promo(p["dur"])
                    if w0:
                        mult = CURVA_TOTTUS.get(semana - w0, 1.0)
                        nota = f"curva Tottus semana +{semana - w0} → {mult:.2f}×"
            prop = base * mult
            row.update({"A": A, "B": B, "base": base, "mult": mult, "prop": prop, "fuente": fuente, "nota": nota,
                        "real_rr": rr1, "si_ant_sem": g(h, semana, "si0"), "si_ant_3s": s0, "yoy": yoy, "estac": estac,
                        "cv": cv, "conf": "baja" if (nuevo or cv is None or cv > 0.8 or base < 1) else ("alta" if cv < 0.4 else "media")})
        elif sub in vr and vr_w:
            real = vr[sub].get(sku, {})
            vals = [real.get(w, 0) for w in vr_w]
            _q = [0.5 ** ((vr_w[-1] - w) / VIDA_MEDIA) for w in vr_w]
            media = sum(v * q for v, q in zip(vals, _q)) / sum(_q)
            cv = statistics.pstdev(vals) / media if media > 0 else None
            prop = media if media > 0 else r["fcst"]
            row.update({"base": prop, "mult": 1.0, "prop": prop, "real_rr": media, "cv": cv,
                        "fuente": f"venta real S{vr_w[0]}–S{vr_w[-1]}" if media > 0 else "sin venta previa (tu FCST)",
                        "nota": "", "conf": "media" if (cv is not None and cv < 0.8) else "baja"})
        else:
            row.update({"base": r["fcst"], "mult": 1.0, "prop": r["fcst"], "fuente": "sin historia (tu FCST)", "nota": "", "conf": "sin historia"})
        row["dif"] = row["prop"] - row["fcst"]
        row["dif_pct"] = row["prop"] / row["fcst"] - 1 if row["fcst"] > 0 else None
        out.append(row)
    if cpfr:
        cp = leer_cpfr(cpfr, {r["sku"] for r in out})
        if tp_path:
            cp["no_normal"] = leer_no_normal(tp_path, set(cp["cerradas"]))
        n = len(cp["cerradas"]) or 1
        for r in out:
            r["cpfr_cadena"] = CPFR_CADENA.get(r["cadena"], r["cadena"])
            v = cp["sku_cad"].get(f'{r["sku"]}|{r["cpfr_cadena"]}')
            if v:
                r["ped_sem"] = v.get("sol", 0) / n          # pedido promedio semanal SKU×cadena CPFR (Cencosud = Jumbo + Santa Isabel)
                r["fcst_cpfr_sem"] = v.get("fcst", 0) / n
    else:
        cp = None
    return {"semana": semana, "anio": anio, "ult_real": ult, "sem_rr": [min(sem_rr), max(sem_rr)], "cpfr": cp,
            "vr_sem": [vr_w[0], vr_w[-1]] if vr_w else None, "filas": out}


def escribir_excel(d, path):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    hdr = PatternFill("solid", fgColor="1D5C8F")
    S = d["semana"]
    # resumen por cadena
    ws = wb.active
    ws.title = f"Resumen S{S}"
    cads = defaultdict(lambda: defaultdict(float))
    for r in d["filas"]:
        c = cads[r["cadena"]]
        c["fcst"] += r["fcst"]; c["prop"] += r["prop"]; c["prev"] += r["fcst_prev"]
        c["real"] += r.get("real_rr") or 0; c["ant"] += r.get("si_ant_sem") or 0
    ws.append(["Cadena", f"Tu FCST S{S} (t)", f"FCST propuesto S{S} (t)", "Propuesto − tu FCST (t)", "Δ%",
               f"Sell In real prom. semanal reciente (t)", f"Sell In S{S} {d['anio'] - 1} (t)", f"Tu FCST prom. S{S-2}–S{S-1} (t)"])
    for k, c in sorted(cads.items(), key=lambda x: -x[1]["fcst"]):
        ws.append([k, round(c["fcst"], 1), round(c["prop"], 1), round(c["prop"] - c["fcst"], 1),
                   c["prop"] / c["fcst"] - 1 if c["fcst"] else None, round(c["real"], 1) or None, round(c["ant"], 1) or None, round(c["prev"], 1)])
        ws.cell(ws.max_row, 5).number_format = "+0%;-0%;0%"
    n = ws.max_row
    ws.append(["Total"] + [f"=SUM({get_column_letter(j)}2:{get_column_letter(j)}{n})" for j in (2, 3, 4)] + [f"=C{n+1}/B{n+1}-1"])
    ws.cell(n + 1, 5).number_format = "+0%;-0%;0%"
    for c in ws[n + 1]:
        c.font = Font(bold=True)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate([30, 13, 15, 15, 9, 18, 15, 16], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    # detalle
    ws = wb.create_sheet(f"SKU x cadena S{S}")
    H = ["SKU", "Producto", "Subcategoría", "Gerencia", "Cadena", f"Tu FCST S{S} (t)", f"FCST propuesto S{S} (t)",
         "Propuesto − tu FCST (t)", "Δ%", "Base (t)", "Método A (t)", "Método B (t)", "Ajuste promo", "Fuente", "Confianza",
         "Sell In real prom. reciente (t)", f"Sell In S{S} {d['anio'] - 1} (t)", "Ritmo año/año", "Estacionalidad", "Promo", "Tu decisión (t)", "Comentario"]
    ws.append(H)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical="center")
    rd = lambda v, k=2: round(v, k) if isinstance(v, (int, float)) else None
    up, dn = PatternFill("solid", fgColor="FBEFDC"), PatternFill("solid", fgColor="FAE3E1")
    for r in sorted(d["filas"], key=lambda r: -abs(r["dif"])):
        if r["fcst"] < 0.01 and r["prop"] < 0.01:
            continue
        ws.append([r["sku"], r["producto"], r["subcat"], r["gerencia"], r["cadena"], rd(r["fcst"]), rd(r["prop"]), rd(r["dif"]),
                   r["dif_pct"], rd(r.get("base")), rd(r.get("A")), rd(r.get("B")), r.get("nota") or None, r["fuente"], r["conf"],
                   rd(r.get("real_rr")), rd(r.get("si_ant_sem")), rd(r.get("yoy")), rd(r.get("estac")), r["promo"] or None, None, None])
        i = ws.max_row
        ws.cell(i, 9).number_format = "+0%;-0%;0%"
        if r["dif_pct"] is not None and abs(r["dif_pct"]) >= 0.2 and abs(r["dif"]) >= 1:
            ws.cell(i, 9).fill = up if r["dif"] < 0 else dn   # tu FCST alto → naranjo (merma); bajo → rojo (quiebre)
        ws.cell(i, 21).fill = PatternFill("solid", fgColor="EEF2F5")
    n = ws.max_row
    ws.append(["Total", None, None, None, None, f"=SUM(F2:F{n})", f"=SUM(G2:G{n})", f"=G{n+1}-F{n+1}", f"=G{n+1}/F{n+1}-1"])
    ws.cell(n + 1, 9).number_format = "+0%;-0%;0%"
    for c in ws[n + 1]:
        c.font = Font(bold=True)
    for j, w in enumerate([11, 40, 28, 12, 26, 11, 12, 12, 8, 10, 10, 10, 22, 22, 11, 12, 12, 9, 9, 24, 12, 26], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:V{n}"
    if d.get("cpfr"):
        cp = d["cpfr"]; sem = "–".join(f"S{w}" for w in cp["cerradas"])
        ws = wb.create_sheet(f"Precisión {sem}")
        ws.append(["Cadena CPFR", f"Tu FCST {sem} (t)", f"Pedido (Solicitado) {sem} (t)", f"Sell In {sem} (t)", "Quebrados (t)", "Sesgo FCST vs pedido",
                   "Venta intermedia + liquidación (t)", "Sesgo FCST vs pedido sin intermedia"])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical="center")
        for c, v in sorted(cp["cadenas"].items(), key=lambda x: -x[1].get("sol", 0)):
            ws.append([c, round(v.get("fcst", 0), 1), round(v.get("sol", 0), 1), round(v.get("si", 0), 1), round(v.get("queb", 0), 1),
                       v["fcst"] / v["sol"] - 1 if v.get("sol") else None])
            nn = cp.get("no_normal", {}).get(c, 0.0)
            ws.cell(ws.max_row, 7).value = round(nn, 1)
            ws.cell(ws.max_row, 8).value = v["fcst"] / (v["sol"] - nn) - 1 if v.get("sol", 0) - nn > 0 else None
            for col in (6, 8):
                ws.cell(ws.max_row, col).number_format = "+0%;-0%;0%"
        for j, w in enumerate([28, 14, 16, 14, 12, 14, 16, 18], 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    m = wb.create_sheet("Cómo se calcula")
    for t in (__doc__ or "").strip().splitlines():
        m.append([t])
    m.append([f"Datos reales hasta S{d['ult_real']}; ritmo reciente S{d['sem_rr'][0]}–S{d['sem_rr'][1]}."])
    m.append(["Δ% en naranjo = tu FCST está sobre la propuesta (riesgo de merma); en rojo = bajo la propuesta (riesgo de quiebre)."])
    m.column_dimensions["A"].width = 130
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--semana", type=int, required=True)
    ap.add_argument("--anio", type=int, default=dt.date.today().year)
    ap.add_argument("--venta-real", default=None)
    ap.add_argument("--cpfr", default=None, help="reporte CPFR semanal (FCST, Solicitado, Sell In por SKU y cadena)")
    ap.add_argument("--out", required=True)
    ap.add_argument("--json", default=None)
    ap.add_argument("--html", default=None, help="página para revisar la semana (usa plantilla_semana.html)")
    a = ap.parse_args()
    d = construir(a.xlsx, a.semana, a.anio, leer_venta_real(a.venta_real) if a.venta_real else None, a.cpfr, a.venta_real)
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    escribir_excel(d, a.out)
    if a.html:
        faltan = json.dumps("".join(f"<div><h3>{t}</h3><p>{d}</p></div>" for t, d in FALTAN), ensure_ascii=False)
        tpl = (Path(__file__).resolve().parent / "plantilla_semana.html").read_text(encoding="utf-8")
        tpl = tpl.replace("/*__DATOS__*/null", json.dumps(d, ensure_ascii=False, default=str)).replace("__S__", str(a.semana)).replace("__FALTAN__", faltan)
        Path(a.html).write_text(tpl, encoding="utf-8")
    if a.json:
        Path(a.json).write_text(json.dumps(d, ensure_ascii=False, default=str))
    f = sum(r["fcst"] for r in d["filas"]); p = sum(r["prop"] for r in d["filas"])
    print(f"OK -> {a.out} | tu FCST S{a.semana} {f:.0f} t | propuesto {p:.0f} t ({p / f - 1:+.1%})")


if __name__ == "__main__":
    main()
